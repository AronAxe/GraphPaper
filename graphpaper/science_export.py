"""Professional APA Word manuscripts and auditable scientific working packages."""
from __future__ import annotations
import csv
import html
import io
import json
import re
import zipfile
from .apa import markdown, title_of, body_of, inline_citations, references, bibliography_sources


def add_inline(paragraph,text):
    for item in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',text):
        run=paragraph.add_run(item.strip('*') if item.startswith('*') else item)
        run.bold=item.startswith('**')
        run.italic=item.startswith('*') and not run.bold


def docx_bytes(project):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document=Document()
    section=document.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
    section.header_distance=section.footer_distance=Inches(.5)
    def typeface(style):
        fonts=style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
            fonts.attrib.pop(qn('w:'+key),None)
        for key in ['ascii','hAnsi','eastAsia','cs']:
            fonts.set(qn('w:'+key),'Times New Roman')
    normal=document.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    typeface(normal)
    normal.paragraph_format.line_spacing=2
    normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.first_line_indent=Inches(.5);normal.paragraph_format.widow_control=True
    for level in range(1,4):
        style=document.styles['Heading '+str(level)]
        style.font.name='Times New Roman';style.font.size=Pt(12);style.font.bold=True
        style.font.italic=level==3;style.font.color.rgb=RGBColor(0,0,0)
        typeface(style)
        style.paragraph_format.line_spacing=2
        style.paragraph_format.space_before=Pt(0);style.paragraph_format.space_after=Pt(0)
        style.paragraph_format.first_line_indent=0;style.paragraph_format.keep_with_next=True
        style.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
    plan=project.research.plan;title=title_of(project)
    header=section.header.paragraphs[0]
    header.paragraph_format.first_line_indent=0;header.paragraph_format.line_spacing=1
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
    header.add_run((plan.running_head or title[:50]).upper()+'\t')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');header._p.append(field)
    def centered(value,bold=False):
        paragraph=document.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent=0;paragraph.add_run(value).bold=bold
        return paragraph
    for _ in range(3):document.add_paragraph().paragraph_format.first_line_indent=0
    centered(title,True);document.add_paragraph().paragraph_format.first_line_indent=0
    centered(plan.author_names or '[Author names required]');centered(plan.affiliation or '[Affiliation required]')
    document.add_paragraph().paragraph_format.first_line_indent=0
    centered('Author Note',True)
    declarations=[plan.author_note]
    for label,value in [('Funding',plan.funding),('Conflicts of interest',plan.conflicts),('Data availability',plan.data_availability),('Ethics',plan.ethics),('AI assistance',plan.ai_disclosure)]:
        if value.strip():declarations.append(label+': '+value)
    for value in declarations:
        if value.strip():document.add_paragraph(value.strip())
    centered('Abstract',True).paragraph_format.page_break_before=True
    paragraph=document.add_paragraph(project.research.abstract);paragraph.paragraph_format.first_line_indent=0
    paragraph=document.add_paragraph();paragraph.add_run('Keywords: ').italic=True
    paragraph.add_run(', '.join(project.research.keywords))
    centered(title,True).paragraph_format.page_break_before=True
    body=inline_citations(project,body_of(project))
    for block in re.split(r'\n\s*\n',body):
        match=re.match(r'^(#{1,4})\s+([^\n]+)(?:\n(.*))?$',block,re.S)
        if match:
            label=match[2].strip()
            if label.lower() not in {'introduction','references','abstract'}:
                document.add_heading(label,min(3,max(1,len(match[1])-1)))
            if match[3]:add_inline(document.add_paragraph(),match[3])
        else:add_inline(document.add_paragraph(),block)
    centered('References',True).paragraph_format.page_break_before=True
    for _,parts in references(project):
        paragraph=document.add_paragraph()
        paragraph.paragraph_format.left_indent=Inches(.5);paragraph.paragraph_format.first_line_indent=Inches(-.5)
        for value,italic in parts:paragraph.add_run(value).italic=italic
    document.core_properties.title=title;document.core_properties.author=plan.author_names
    buffer=io.BytesIO();document.save(buffer)
    return buffer.getvalue()


def bibliography(project,kind='bib'):
    result=[]
    for source in bibliography_sources(project):
        meta=source.scholarly
        if kind=='ris':
            lines=['TY  - '+('UNPB' if meta.preprint else 'JOUR'),'TI  - '+meta.title]
            lines += ['AU  - '+(a.literal or a.family+', '+a.given) for a in meta.authors]
            for key,value in [('PY',meta.year),('JO',meta.journal),('VL',meta.volume),('IS',meta.issue),('SP',meta.pages),('DO',meta.doi),('UR',meta.url)]:
                if value:lines.append(key+'  - '+str(value).replace('\n',' '))
            result.append('\n'.join(lines+['ER  -','']))
        else:
            values={'title':meta.title,'author':' and '.join(a.literal or a.family+', '+a.given for a in meta.authors),'year':meta.year,'journal':meta.journal,'volume':meta.volume,'number':meta.issue,'pages':meta.pages,'doi':meta.doi,'url':meta.url}
            def safe(value):
                return str(value).replace('{','').replace('}','').replace('\n',' ')
            content=',\n'.join('  '+key+' = {'+safe(value)+'}' for key,value in values.items() if value)
            result.append('@'+('misc' if meta.preprint else 'article')+'{'+source.id+',\n'+content+'\n}')
    return '\n\n'.join(result)+'\n'


def evidence_csv(project):
    output=io.StringIO(newline='');writer=csv.writer(output)
    writer.writerow(['ID','Title','Year','DOI','Database provenance','Decision','Reason','Content scope','Design','Sample','Findings','Limitations','Exact source quotations'])
    def safe(value):
        value=str(value or '')
        return "'"+value if value.startswith(('=','+','-','@','\t','\r')) else value
    for record in project.research.records:
        meta=record.metadata;appraisal=record.appraisal
        values=[record.source_id,meta.title,meta.year,meta.doi,', '.join(meta.metadata_sources),record.decision,record.reason,meta.content_scope,appraisal.get('design'),appraisal.get('sample'),appraisal.get('findings'),appraisal.get('limitations'),'; '.join(str(c.get('quote','')) for c in appraisal.get('claims',[]) if c.get('exact_match'))]
        writer.writerow([safe(v) for v in values])
    return '\ufeff'+output.getvalue()


def export_science(project,kind):
    if kind in {'md','apa-md'}:return markdown(project).encode(),'text/markdown; charset=utf-8','.md'
    if kind in {'docx','apa-docx'}:return docx_bytes(project),'application/vnd.openxmlformats-officedocument.wordprocessingml.document','.docx'
    if kind in {'bib','ris'}:return bibliography(project,kind).encode(),'text/plain; charset=utf-8','.'+kind
    if kind=='evidence':return evidence_csv(project).encode(),'text/csv; charset=utf-8','.csv'
    if kind=='search-log':return json.dumps(project.research.searches,ensure_ascii=False,indent=2).encode(),'application/json','.json'
    if kind=='html':
        from .export import simple_html
        content='<!doctype html><html><meta charset="utf-8"><title>'+html.escape(title_of(project))+'</title><style>body{max-width:6.5in;margin:1in auto;font:12pt/2 "Times New Roman",serif}p{margin:0;text-indent:.5in}h1,h2,h3{font-size:12pt;margin:0}h1,h2{text-align:center}@page{size:letter;margin:1in}@media print{body{margin:0}h1,h2{break-after:avoid}}</style><body>'+simple_html(markdown(project))+'</body></html>'
        return content.encode(),'text/html; charset=utf-8','.html'
    if kind in {'research-package','submission'}:
        from .science import readiness
        status=readiness(project)
        if kind=='submission' and not status['submission_allowed']:
            raise ValueError('Complete the submission checklist and author confirmations first. A working research package can still be exported.')
        memory=io.BytesIO()
        with zipfile.ZipFile(memory,'w',zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('manuscript.docx',docx_bytes(project));archive.writestr('manuscript.md',markdown(project))
            archive.writestr('references.bib',bibliography(project));archive.writestr('references.ris',bibliography(project,'ris'))
            archive.writestr('evidence.csv',evidence_csv(project));archive.writestr('search-log.json',json.dumps(project.research.searches,ensure_ascii=False,indent=2))
            archive.writestr('protocol.json',project.research.plan.model_dump_json(indent=2))
            archive.writestr('submission-checks.json',json.dumps(status,indent=2))
            archive.writestr('README.txt','GraphPaper scientific research package. '+('Author confirmations recorded. ' if kind=='submission' else 'Working export; unresolved checks may remain. ')+'Formatting and software checks do not constitute journal acceptance or ethics approval. Full source manuscripts and credentials are not included.\n')
        return memory.getvalue(),'application/zip','.zip'
    raise ValueError('Unknown scientific export format')
