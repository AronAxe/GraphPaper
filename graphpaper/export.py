from __future__ import annotations
import html
import io
import json
import re

from .models import Project


def referenced_sources(p: Project):
    refs = set(re.findall(r"\[(S\d+)\]", p.draft))
    return [s for s in p.sources if s.id in refs and s.enabled and s.role == "evidence"]


def markdown(p: Project):
    text = p.draft
    sources = referenced_sources(p) if p.mode in {"nonfiction", "polemic"} else []
    if sources:
        text += "\n\n---\n\n## Sources\n\n"
        for s in sources:
            text += f"[{s.id}] {s.title}" + (f" — {s.author}" if s.author else "") + (f" ({s.published})" if s.published else "") + (f"\n{s.url}" if s.url else "") + "\n\n"
    return text


def simple_html(text):
    """Small safe Markdown subset, never allows raw HTML or remote assets."""
    out = []
    for block in re.split(r"\n\s*\n", text.strip()):
        t = html.escape(block)
        t = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", t)
        t = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<em>\1</em>", t)
        if re.match(r"^#{1,3} ", block):
            n = len(block) - len(block.lstrip("#"))
            out.append(f"<h{n}>{t[n+1:]}</h{n}>")
        elif block.strip() in {"---", "* * *"}:
            out.append("<hr>")
        else:
            out.append("<p>" + t.replace("\n", "<br>") + "</p>")
    return "\n".join(out)


def export(p: Project, kind: str):
    if p.mode == 'science' and kind not in {'json','graph'}:
        from .science_export import export_science
        return export_science(p,kind)
    if kind == "json":
        return p.model_dump_json(indent=2).encode(), "application/json", ".graphpaper.json"
    if kind == "graph":
        return p.graph.model_dump_json(indent=2).encode(), "application/json", ".graph.json"
    if kind == "md":
        return markdown(p).encode(), "text/markdown; charset=utf-8", ".md"
    if kind == "html":
        text = '<!doctype html><html lang="en"><meta charset="utf-8"><title>' + html.escape(p.title) + '</title><style>body{max-width:740px;margin:64px auto;padding:0 28px;font:18px/1.75 Georgia,serif;color:#26282b}h1{font-size:2.3em;line-height:1.16}h2{margin-top:2em}p{overflow-wrap:anywhere}hr{border:0;border-top:1px solid #ddd;margin:2em 0}@media print{body{margin:0;max-width:none;font-size:11pt}h1,h2,h3{break-after:avoid}p{orphans:3;widows:3}@page{size:A4;margin:24mm}}</style><body>' + simple_html(markdown(p)) + '</body></html>'
        return text.encode(), "text/html; charset=utf-8", ".html"
    if kind == "docx":
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
        sec.top_margin = sec.bottom_margin = Inches(0.9)
        sec.left_margin = sec.right_margin = Inches(0.95)
        style = doc.styles["Normal"]
        style.font.name, style.font.size = "Calibri", Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.18
        doc.core_properties.title, doc.core_properties.author = p.title, ""
        for block in re.split(r"\n\s*\n", markdown(p).strip()):
            m = re.match(r"^(#{1,3}) (.*)", block, re.S)
            if m:
                para = doc.add_heading(m[2], len(m[1]))
                for r in para.runs:
                    r.font.color.rgb = RGBColor.from_string("243E38")
            elif block.strip() in {"---", "* * *"}:
                doc.add_paragraph("* * *")
            else:
                para = doc.add_paragraph()
                for token in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", block):
                    run = para.add_run(token.strip("*") if token.startswith("*") else token)
                    run.bold = token.startswith("**")
                    run.italic = token.startswith("*") and not run.bold
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"
    raise ValueError("Unknown export format")
