import io
from docx import Document
from graphpaper.science_export import docx_bytes
from graphpaper.apa import inline_citations,reference_parts,citation_map
from graphpaper.science import require_evidence,record_source
from graphpaper.models import Project,Source
from .test_scholarly_research import prepared
from .science_fixtures import papers
import pytest


def test_apa_reference_follows_clause_and_precedes_period():
    p=prepared()
    assert inline_citations(p,'The result was uncertain. [S1] [S2]')=='The result was uncertain (Example & Researcher, 2024, 2025).'


def test_major_sections_use_page_break_before_not_empty_break_paragraphs():
    document=Document(io.BytesIO(docx_bytes(prepared())))
    assert sum(bool(p.paragraph_format.page_break_before) for p in document.paragraphs)==3
    assert '<w:br w:type="page"' not in document.element.xml


def test_apa_heading_fonts_do_not_inherit_inconsistent_office_theme():
    document=Document(io.BytesIO(docx_bytes(prepared())))
    for name in ['Normal','Heading 1','Heading 2','Heading 3']:
        xml=document.styles[name].element.xml
        assert 'asciiTheme' not in xml and 'hAnsiTheme' not in xml
        assert 'Times New Roman' in xml


def test_specialized_reference_override_preserves_author_intent():
    p=prepared();s=p.sources[0];s.scholarly.reference_override='Example, A. (2024). *A book title*. Publisher.'
    parts=reference_parts(s,citation_map(p))
    assert ''.join(text for text,_ in parts)== 'Example, A. (2024). A book title. Publisher.'
    assert ('A book title',True) in parts


def test_empirical_mode_requires_completed_author_results():
    p=prepared();p.research.plan.article_type='empirical'
    with pytest.raises(ValueError,match='requires your own'):require_evidence(p)
    source=Source(id='S3',title='Completed experiment report',text='A synthetic methods/results report.',role='evidence')
    p.sources.append(source);p.research.plan.empirical_results_source_ids=['S3']
    assert len(require_evidence(p))==2


def test_cannot_attach_one_source_to_two_scholarly_papers(client,app):
    p=Project(mode='science');p.research.records=papers()
    for r in p.research.records:r.decision='include';record_source(p,r)
    app.state.store.create(p)
    response=client.post('/api/projects/'+p.id+'/research/records/'+p.research.records[1].id+'/attach',json={'version':p.version,'source_id':p.research.records[0].source_id,'confirm_identity':True,'full_text':True})
    assert response.status_code==400 and 'different paper' in response.json()['detail']


def test_reference_edit_cannot_silently_clear_retraction_flag(client,app):
    p=Project(mode='science');p.research.records=papers();r=p.research.records[0];r.metadata.retracted=True
    app.state.store.create(p);meta=r.metadata.model_dump();meta['retracted']=False
    response=client.patch('/api/projects/'+p.id+'/research/records/'+r.id,json={'version':0,'metadata':meta})
    assert response.status_code==200 and response.json()['research']['records'][0]['metadata']['retracted']
