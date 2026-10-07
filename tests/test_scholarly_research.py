import io
import json
import time
import zipfile
from pathlib import Path
import httpx
import pytest
from defusedxml.common import DefusedXmlException
from graphpaper.models import Project,Source,Settings
from graphpaper.science_models import ResearchPlan,ResearchAuthor
from graphpaper.scholarly import parse_pubmed,parse_semantic,parse_arxiv,parse_crossref,parse_europe,merge_records,ScholarlyClient,doi
from graphpaper.science import record_source,readiness,research_fingerprint,search_literature
from graphpaper.apa import authors_reference,inline_citations,citation_map,markdown
from graphpaper.science_export import docx_bytes,export_science,evidence_csv
from graphpaper.ingest import digest
from graphpaper.pipeline import Job
from .science_fixtures import papers,FakeScholar,ScienceClients

PUBMED=b'''<PubmedArticleSet><PubmedArticle><MedlineCitation><PMID>123</PMID><Article><ArticleTitle>Test paper with <i>markup</i></ArticleTitle><Journal><Title>Journal of Testing</Title><JournalIssue><Volume>3</Volume><Issue>2</Issue><PubDate><Year>2024</Year></PubDate></JournalIssue></Journal><Pagination><MedlinePgn>21-28</MedlinePgn></Pagination><Abstract><AbstractText Label="RESULTS">The result was uncertain.</AbstractText></Abstract><AuthorList><Author><LastName>van Example</LastName><ForeName>Anna Bea</ForeName></Author></AuthorList><PublicationTypeList><PublicationType>Journal Article</PublicationType></PublicationTypeList></Article></MedlineCitation><PubmedData><ArticleIdList><ArticleId IdType="doi">10.9999/test</ArticleId><ArticleId IdType="pmc">PMC456</ArticleId></ArticleIdList></PubmedData></PubmedArticle></PubmedArticleSet>'''


def test_pubmed_metadata_abstract_and_author_names():
    r=parse_pubmed(PUBMED)[0]
    assert r.metadata.title=='Test paper with markup'
    assert r.metadata.authors[0].family=='van Example'
    assert r.metadata.doi=='10.9999/test' and r.metadata.pmcid=='PMC456'
    assert r.metadata.content_scope=='abstract' and 'uncertain' in r.abstract


def test_pubmed_missing_abstract_is_metadata_only():
    r=parse_pubmed(PUBMED.replace(b'<Abstract><AbstractText Label="RESULTS">The result was uncertain.</AbstractText></Abstract>',b''))[0]
    assert not r.abstract and r.metadata.content_scope=='metadata'


def test_pubmed_retraction_flag():
    assert parse_pubmed(PUBMED.replace(b'Journal Article',b'Retracted Publication'))[0].metadata.retracted


def test_xml_external_entities_not_resolved():
    with pytest.raises(DefusedXmlException):parse_pubmed(b'<!DOCTYPE x [<!ENTITY steal SYSTEM "file:///not-an-allowed-input">]><x>&steal;</x>')


def test_semantic_external_ids_and_access_scope():
    rows=parse_semantic({'data':[{'title':'Title','year':2024,'authors':[{'name':'A. Tester'}],'abstract':'Abstract.','externalIds':{'DOI':'10.9999/test','PubMed':'123'},'openAccessPdf':{'url':'https://example.org/a.pdf'},'journal':{'name':'Journal','volume':'3','pages':'1-9'}}]})
    assert rows[0].metadata.doi=='10.9999/test'
    assert rows[0].metadata.content_scope=='abstract'
    assert rows[0].metadata.fulltext_url.endswith('.pdf')


def test_arxiv_atom_remains_labelled_preprint():
    body=b'''<feed xmlns="http://www.w3.org/2005/Atom" xmlns:o="http://a9.com/-/spec/opensearch/1.1/"><o:totalResults>90</o:totalResults><entry><id>http://arxiv.org/abs/2401.00001v2</id><title> Test title </title><summary>Abstract text</summary><published>2024-01-01</published><author><name>Bea Tester</name></author><link title="pdf" href="http://arxiv.org/pdf/2401.00001v2" /></entry></feed>'''
    rows,total=parse_arxiv(body)
    assert total==90 and rows[0].metadata.preprint
    assert rows[0].metadata.arxiv_id=='2401.00001v2'


def test_crossref_structured_bibliography():
    raw={'message':{'items':[{'title':['Test title'],'DOI':'10.9999/test','author':[{'family':'Tester','given':'Alex'}],'container-title':['Journal'],'published':{'date-parts':[[2023,2,1]]},'type':'journal-article','abstract':'<p>Finding.</p>','volume':'2','issue':'3','page':'1-8'}]}}
    r=parse_crossref(raw)[0]
    assert r.metadata.year==2023 and r.abstract=='Finding.'
    assert r.metadata.authors[0].family=='Tester' and not r.metadata.preprint


def test_europe_pmc_open_access_ids():
    raw={'resultList':{'result':[{'title':'Paper','id':'123','source':'MED','pmid':'123','pmcid':'PMC456','pubYear':'2024','abstractText':'Summary','authorList':{'author':[{'lastName':'Example','firstName':'A'}]},'journalInfo':{'journal':{'title':'Journal'}}}]}}
    r=parse_europe(raw)[0]
    assert r.metadata.pmcid=='PMC456' and r.metadata.content_scope=='abstract'


def test_duplicate_merging_preserves_screening_and_provenance():
    first=papers();first[0].decision='exclude';first[0].reason='Not eligible'
    second=papers();second[0].metadata.metadata_sources=['crossref']
    result,count=merge_records(first,second)
    assert len(result)==2 and count==2
    assert result[0].decision=='exclude' and result[0].reason=='Not eligible'
    assert set(result[0].metadata.metadata_sources)=={'pubmed','crossref'}


@pytest.mark.parametrize('value,expected',[('https://doi.org/10.9999/ABC','10.9999/abc'),('doi: 10.9999/test','10.9999/test'),('not-a-doi','')])
def test_doi_normalization(value,expected):assert doi(value)==expected


def test_research_dates_are_validated():
    with pytest.raises(ValueError):ResearchPlan(year_from=2025,year_to=2020)


def test_search_transport_uses_correct_eutils_endpoints():
    requests=[]
    def handler(request):
        requests.append(request)
        if request.url.path.endswith('esearch.fcgi'):return httpx.Response(200,json={'esearchresult':{'count':'4','idlist':['123']}})
        return httpx.Response(200,content=PUBMED)
    c=ScholarlyClient(transport=httpx.MockTransport(handler))
    records,log=c.search('pubmed','test query',ResearchPlan(per_database=1,year_from=2020))
    assert len(records)==1 and log['truncated'] and log['total_hits']==4
    assert 'Date - Publication' in log['query']
    assert len(requests)==2


def test_partial_database_failures_are_logged_not_empty_success(tmp_path):
    class Partial(FakeScholar):
        def search(self,db,query,plan):
            if db=='semantic_scholar':raise ValueError('Semantic Scholar returned HTTP 429')
            return super().search(db,query,plan)
    c=ScienceClients();c.scholarly_factory=Partial
    p=Project(mode='science');p.research.plan=ResearchPlan(question='Memory',databases=['pubmed','semantic_scholar'])
    search_literature(p,c,None,Job(p.id,'science-search'))
    assert len(p.research.records)==2
    assert [x['status'] for x in p.research.searches]==['ok','failed']
    assert '429' in p.research.searches[-1]['error']


def prepared():
    p=Project(mode='science',title='A test manuscript')
    p.research.records=papers()
    for r in p.research.records:r.decision='include';record_source(p,r)
    p.draft='# A test manuscript\n\n## Introduction\n\nThe result was uncertain. [S1] [S2]'
    p.research.abstract='An abstract of this synthetic workflow.'
    p.research.plan=ResearchPlan(question='Test',author_names='Alex Example',affiliation='Test Institute',funding='None',conflicts='None',data_availability='Synthetic material',ethics='Not applicable')
    p.research.searches=[{'database':'pubmed','query':'test','status':'ok','truncated':False}]
    p.review.draft_hash=digest(p.draft)
    return p


def test_apa_author_counts_and_initials():
    authors=[ResearchAuthor(family=f'Tester{i}',given='Alex Bea') for i in range(21)]
    value=authors_reference(authors)
    assert 'Tester18' in value and 'Tester19' not in value and 'Tester20' in value and '. . .' in value
    assert authors_reference(authors[:2])=='Tester0, A. B., & Tester1, A. B.'


def test_apa_same_year_disambiguation_and_adjacent_citations():
    p=prepared();p.sources[1].scholarly.year=2024
    mapping=citation_map(p)
    assert {mapping['S1']['year'],mapping['S2']['year']}=={'2024a','2024b'}
    value=inline_citations(p,'Evidence [S2] [S1].')
    assert value=='Evidence (Example & Researcher, 2024a, 2024b).'
    assert 'UNRESOLVED' in inline_citations(p,'[S99]')


def test_apa_word_dimensions_styles_header_and_references():
    from docx import Document
    from docx.shared import Inches
    p=prepared();doc=Document(io.BytesIO(docx_bytes(p)))
    assert doc.sections[0].page_width==Inches(8.5)
    assert doc.sections[0].top_margin==Inches(1)
    assert doc.styles['Normal'].paragraph_format.line_spacing==2
    assert doc.styles['Normal'].font.name=='Times New Roman'
    assert 'PAGE' in doc.sections[0].header._element.xml
    paragraphs=[x for x in doc.paragraphs if 'https://doi.org' in x.text]
    assert len(paragraphs)==2 and paragraphs[0].paragraph_format.first_line_indent==Inches(-.5)
    assert any(r.italic and 'Test Fixture Journal' in r.text for r in paragraphs[0].runs)
    assert '[S1]' not in '\n'.join(x.text for x in doc.paragraphs)


def test_submission_is_blocked_until_current_author_checks():
    p=prepared()
    with pytest.raises(ValueError):export_science(p,'submission')
    p.research.acknowledgements={k:True for k in readiness(p)['author_checks']}
    p.research.confirmation_hash=digest(research_fingerprint(p)+p.draft+p.research.abstract)
    assert readiness(p)['submission_allowed']
    data,_,_=export_science(p,'submission')
    with zipfile.ZipFile(io.BytesIO(data)) as z:assert {'manuscript.docx','references.bib','search-log.json','evidence.csv'}<=set(z.namelist())
    p.draft+=' Changed.'
    assert not readiness(p)['submission_allowed']


def test_systematic_review_cannot_hide_capped_search():
    p=prepared();p.research.plan.article_type='systematic_review';p.research.searches[0]['truncated']=True
    assert any('incomplete' in b.lower() for b in readiness(p)['blockers'])


def test_csv_formula_safety():
    p=prepared();p.research.records[0].metadata.title='=1+1'
    assert "'=1+1" in evidence_csv(p)


def wait_job(client,jid):
    for _ in range(250):
        j=client.get('/api/jobs/'+jid).json()
        if j['state'] not in {'queued','running'}:
            assert j['state']=='completed',j
            return j
        time.sleep(.02)
    raise AssertionError('Job did not finish')


def test_science_project_full_workflow_and_export(client,app):
    app.state.runner.clients_factory=ScienceClients
    app.state.store.set_settings(Settings(model='test',allow_cloud=True,refine=False).model_dump())
    p=client.post('/api/projects',json={'title':'Science test','mode':'science'}).json();base='/api/projects/'+p['id']
    assert p['mode']=='science' and p['brief']['format']=='APA scientific manuscript'
    p=client.put(base+'/research/plan',json={'version':p['version'],'plan':{'question':'Memory and sleep','databases':['pubmed']}}).json()
    def run(action):wait_job(client,client.post(base+'/jobs',json={'action':action}).json()['id'])
    run('science-search');p=client.get(base).json();assert len(p['research']['records'])==2
    p=client.post(base+'/research/screen',json={'version':p['version'],'ids':[r['id'] for r in p['research']['records']],'decision':'include'}).json()
    run('science-appraise');run('science-outline');run('science-draft')
    p=client.get(base).json();assert p['research']['abstract'] and '[S1]' in p['draft']
    assert p['research']['records'][0]['appraisal']['claims'][0]['exact_match']
    preview=client.get(base+'/research/preview').json()['markdown']
    assert '(Example & Researcher, 2024,' in preview and '## References' in preview
    for kind in ['docx','md','bib','ris','evidence','search-log','research-package']:
        response=client.get(base+'/export/'+kind);assert response.status_code==200,(kind,response.text[:400])
    assert client.get(base+'/export/submission').status_code==400


def test_screening_and_project_boundaries(client):
    p=client.post('/api/projects',json={'mode':'science'}).json()
    result=client.put('/api/projects/'+p['id']+'/research/plan',json={'version':99,'plan':{}})
    assert result.status_code==409
    n=client.post('/api/projects',json={'mode':'nonfiction'}).json()
    assert client.put('/api/projects/'+n['id']+'/research/plan',json={'version':0,'plan':{}}).status_code==400
